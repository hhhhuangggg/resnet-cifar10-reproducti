"""Build a clean online-document-style DOCX from the verified ResNet report."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT_ROOT / "docs" / "ResNet复现报告.md"
OUTPUT = PROJECT_ROOT / "ResNet复现报告-飞书版-未清理.docx"

ASCII_FONT = "Arial"
CJK_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"
BLACK = "000000"
MUTED = "555555"
BORDER = "DADCE0"
CODE_FILL = "F8F9FA"
CONTENT_DXA = 9360


def set_run_font(
    run,
    *,
    size: float = 11,
    bold: bool | None = None,
    color: str = BLACK,
    mono: bool = False,
) -> None:
    font = MONO_FONT if mono else ASCII_FONT
    run.font.name = font
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    fonts.set(qn("w:eastAsia"), MONO_FONT if mono else CJK_FONT)


def set_paragraph_spacing(
    paragraph,
    *,
    before: float = 0,
    after: float = 8,
    line: float = 1.15,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"table widths must total {CONTENT_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def column_widths(count: int) -> list[int]:
    patterns = {
        2: [3000, 6360],
        3: [2400, 3480, 3480],
        4: [2100, 2420, 2420, 2420],
        5: [1750, 1902, 1902, 1902, 1904],
        6: [1450, 1582, 1582, 1582, 1582, 1582],
    }
    if count in patterns:
        return patterns[count]
    base, remainder = divmod(CONTENT_DXA, count)
    return [base + (1 if index < remainder else 0) for index in range(count)]


def add_markdown_runs(paragraph, text: str, *, size: float = 11, color: str = BLACK) -> None:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`|\*[^*]+?\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size - 0.5, color=color, mono=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color)
            run.italic = True
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, color=color)


def add_numbering(document: Document) -> tuple[int, int]:
    numbering = document.part.numbering_part.element
    existing_abstract = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    existing_num = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    next_abstract = max(existing_abstract or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def create(fmt: str, text: str, abstract_id: int, num_id: int) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        level.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        level.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        level.append(p_pr)
        abstract.append(level)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    create("bullet", "●", next_abstract, next_num)
    create("decimal", "%1.", next_abstract + 1, next_num + 1)
    return next_num, next_num + 1


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), ASCII_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        "Heading 1": (20, BLACK, 20, 6),
        "Heading 2": (16, BLACK, 18, 6),
        "Heading 3": (14, "434343", 16, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = ASCII_FONT
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), ASCII_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), ASCII_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    set_table_geometry(table, column_widths(len(rows[0])))
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if col_index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            set_paragraph_spacing(paragraph, after=0, line=1.15)
            add_markdown_runs(
                paragraph,
                value,
                size=9.5,
            )
            for run in paragraph.runs:
                run.bold = row_index == 0
        if row_index == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
    spacer = document.add_paragraph()
    set_paragraph_spacing(spacer, after=4)


def add_figure(document: Document, path: Path, caption: str, width: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=4, after=4)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(cap, after=10)
        run = cap.add_run(caption)
        set_run_font(run, size=9, color=MUTED)


def build() -> None:
    document = Document()
    configure_document(document)
    bullet_num, decimal_num = add_numbering(document)

    title = document.add_paragraph()
    set_paragraph_spacing(title, before=0, after=3)
    title_run = title.add_run("ResNet 方法实现与 CIFAR-10 复现报告")
    set_run_font(title_run, size=26, bold=False)

    subtitle = document.add_paragraph()
    set_paragraph_spacing(subtitle, after=8)
    subtitle_run = subtitle.add_run(
        "从 PlainNet baseline 到 64k 正式训练与私有 GitHub 复现"
    )
    set_run_font(subtitle_run, size=13, color="434343")

    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(source_lines):
        raw = source_lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                paragraph = document.add_paragraph()
                set_paragraph_spacing(paragraph, before=2, after=8, line=1.1)
                p_pr = paragraph._p.get_or_add_pPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), CODE_FILL)
                p_pr.append(shading)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, size=9.5, color="202124", mono=True)
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(source_lines):
            table_lines: list[str] = []
            while index < len(source_lines) and source_lines[index].strip().startswith("|"):
                table_lines.append(source_lines[index].strip())
                index += 1
            parsed = [
                [cell.strip() for cell in line.strip("|").split("|")]
                for line in table_lines
            ]
            if len(parsed) >= 2 and all(
                re.fullmatch(r":?-{3,}:?", cell) for cell in parsed[1]
            ):
                parsed.pop(1)
            add_table(document, parsed)
            continue

        if not stripped:
            index += 1
            continue
        if stripped == "<!-- pagebreak -->":
            document.add_page_break()
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = SOURCE.parent / Path(image_match.group(2))
            if not image_path.exists():
                raise FileNotFoundError(f"Markdown image not found: {image_path}")
            add_figure(document, image_path, "", 6.25)
        elif stripped.startswith("# "):
            pass
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
        elif stripped.startswith("#### "):
            document.add_heading(stripped[5:], level=3)
        elif re.match(r"^-\s+", stripped):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            set_paragraph_spacing(paragraph, after=4)
            add_markdown_runs(
                paragraph,
                "• " + re.sub(r"^-\s+", "", stripped),
            )
        elif re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            set_paragraph_spacing(paragraph, after=4)
            add_markdown_runs(paragraph, stripped)
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            set_paragraph_spacing(paragraph, before=4, after=8)
            add_markdown_runs(
                paragraph,
                stripped.lstrip("> ").strip(),
                color="434343",
            )
        else:
            paragraph = document.add_paragraph()
            set_paragraph_spacing(paragraph)
            add_markdown_runs(paragraph, stripped)
        index += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
