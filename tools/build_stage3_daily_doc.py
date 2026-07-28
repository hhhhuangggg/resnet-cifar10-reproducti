from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_stage2_daily_doc import (
    MID_GRAY,
    add_body,
    add_code,
    add_page_number,
    add_table,
    configure_styles,
    set_run_font,
)


OUTPUT = Path("outputs/2026-07-28 ResNet复现学习日报.docx")


def add_labeled_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(f"{label}："), bold=True, color="1F4D78")
    set_run_font(paragraph.add_run(text))


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("ResNet 论文复现 · 第三阶段日报"), size=9, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("ResNet 复现学习日报"), size=23, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    set_run_font(
        subtitle.add_run("第三阶段：CIFAR-10 PlainNet 网络结构实现与验证"),
        size=12,
        color=MID_GRAY,
    )

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["日期", "2026年7月28日"],
            ["项目目录", r"E:\文档\resnet复现"],
            ["本阶段目标", "实现并验证 PlainNet-20/32/44/56，不加载数据集、不训练"],
            ["最终状态", "第三阶段完成；完整项目 37 项自动测试通过"],
        ],
        [1800, 7560],
    )

    doc.add_heading("一、今日工作概述", level=1)
    add_body(
        doc,
        "本阶段从空的 PlainNet 模型文件开始，按照“先写测试、确认失败、补充最小实现、"
        "再次验证”的顺序，逐步完成 PlainBlock、完整 PlainNet、模型注册、结构测试、"
        "参数量统计和 GPU 前向传播。整个过程只验证网络结构，没有下载 CIFAR-10，"
        "也没有计算 loss、反向传播或更新模型参数。"
    )
    add_labeled_paragraph(
        doc,
        "核心成果",
        "PlainNet-20、32、44、56 均可通过 create_model() 创建，"
        "输入 [N,3,32,32] 后输出 [N,10]，网络深度、阶段尺寸和参数量均通过验证。",
    )
    add_labeled_paragraph(
        doc,
        "学习方式",
        "没有一次性复制完整模型，而是把网络拆成最小可验证单元，亲自观察测试从失败到通过。",
    )

    doc.add_heading("二、第三阶段完成的功能", level=1)
    add_table(
        doc,
        ["模块", "完成内容", "验证方式"],
        [
            ["PlainBlock", "两组 Conv-BN-ReLU，无残差连接", "普通形状、降采样及非法参数测试"],
            ["PlainNet stem", "3通道 RGB 输入转换为16通道特征", "[2,3,32,32] → [2,16,32,32]"],
            ["三个阶段", "按 n 重复 PlainBlock，并在阶段交界降采样", "模块数量和中间形状测试"],
            ["分类头", "全局平均池化、展平、全连接层", "完整输出 [2,10]"],
            ["初始化", "卷积 Kaiming；BN weight=1、bias=0", "调用测试和实际参数测试"],
            ["模型入口", "plain20/32/44/56 名称映射到 PlainNet", "四种型号创建与深度测试"],
            ["GPU验证", "PlainNet-20 在 RTX 3060 上前向传播", "输出 [1,10]，设备 cuda:0"],
        ],
        [1800, 4500, 3060],
    )

    doc.add_heading("三、PlainBlock 的原理与实现", level=1)
    add_code(
        doc,
        "输入 x\n"
        "→ Conv2d(3×3, stride可为1或2, padding=1, bias=False)\n"
        "→ BatchNorm2d → ReLU\n"
        "→ Conv2d(3×3, stride=1, padding=1, bias=False)\n"
        "→ BatchNorm2d → ReLU\n"
        "→ 输出 F(x)"
    )
    add_body(
        doc,
        "PlainBlock 只计算 F(x)，不包含 shortcut，也不执行 F(x)+x。"
        "第一层卷积接收调用者传入的 stride，因此既能保持尺寸，也能在阶段交界处完成降采样；"
        "第二层卷积固定使用 stride=1，防止同一个模块连续缩小两次。"
    )
    add_table(
        doc,
        ["模块用法", "输入", "输出", "作用"],
        [
            ["PlainBlock(16,16,1)", "[2,16,32,32]", "[2,16,32,32]", "阶段内部保持形状"],
            ["PlainBlock(16,32,2)", "[2,16,32,32]", "[2,32,16,16]", "通道翻倍、尺寸减半"],
        ],
        [2300, 2300, 2300, 2460],
    )
    add_labeled_paragraph(
        doc,
        "卷积不使用 bias",
        "卷积后立即接 BatchNorm，BatchNorm 已包含可学习平移参数，额外卷积偏置没有必要。",
    )
    add_labeled_paragraph(
        doc,
        "参数校验",
        "in_channels、out_channels 和 stride 必须是大于0的整数，非法值在创建模块时立即报错。",
    )

    doc.add_heading("四、完整 PlainNet 架构", level=1)
    add_code(
        doc,
        "输入 [N,3,32,32]\n"
        "→ stem [N,16,32,32]\n"
        "→ stage1 [N,16,32,32]\n"
        "→ stage2 [N,32,16,16]\n"
        "→ stage3 [N,64,8,8]\n"
        "→ AdaptiveAvgPool2d [N,64,1,1]\n"
        "→ flatten [N,64]\n"
        "→ Linear [N,num_classes]"
    )
    add_table(
        doc,
        ["阶段", "输出通道", "空间尺寸", "首模块 stride"],
        [
            ["stem", "16", "32×32", "1"],
            ["stage1", "16", "32×32", "1"],
            ["stage2", "32", "16×16", "2"],
            ["stage3", "64", "8×8", "2"],
        ],
        [2200, 2200, 2500, 2460],
    )
    add_body(
        doc,
        "每个阶段的第一个 PlainBlock 单独创建，用来处理通道变化和可能的降采样；"
        "剩余 n-1 个模块保持当前通道数并固定 stride=1。self.in_channels 在构造阶段依次记录"
        "前一阶段的输出通道，变化为 16→16→32→64。"
    )

    doc.add_heading("五、_make_stage() 的关键作用", level=1)
    add_body(
        doc,
        "_make_stage() 避免手工逐层编写几十个 PlainBlock。它先创建阶段首模块，"
        "随后更新 self.in_channels，再用循环创建剩余模块，最后通过 nn.Sequential(*blocks)"
        "把列表中的模块按顺序登记为一个阶段。"
    )
    add_code(
        doc,
        "第一个模块：PlainBlock(当前通道, 目标通道, first_stride)\n"
        "更新当前通道：self.in_channels = out_channels\n"
        "剩余模块：PlainBlock(目标通道, 目标通道, stride=1)"
    )
    add_labeled_paragraph(
        doc,
        "本步骤的理解",
        "变量 n 不只是一个标签；测试通过 len(stage1/2/3) 直接确认每个阶段实际创建了 n 个模块。",
    )

    doc.add_heading("六、深度公式与真实层数", level=1)
    add_body(
        doc,
        "CIFAR PlainNet 的标称深度满足 depth=6n+2。开头 stem 有1个卷积层；"
        "三个阶段各有 n 个 PlainBlock，每个 PlainBlock 含2个卷积层；最后有1个全连接层。"
        "因此总深度为 1+3×n×2+1。"
    )
    add_table(
        doc,
        ["模型", "n", "PlainBlock总数", "Conv2d", "Linear", "真实深度"],
        [
            ["PlainNet-20", "3", "9", "19", "1", "20"],
            ["PlainNet-32", "5", "15", "31", "1", "32"],
            ["PlainNet-44", "7", "21", "43", "1", "44"],
            ["PlainNet-56", "9", "27", "55", "1", "56"],
        ],
        [2100, 900, 1750, 1500, 1400, 1710],
    )
    add_body(
        doc,
        "测试没有只读取 model.depth，而是遍历模型并统计 Conv2d 与 Linear 的真实数量。"
        "BatchNorm、ReLU、池化和容器模块不计入论文这里所说的网络深度。"
    )

    doc.add_heading("七、权重初始化", level=1)
    add_table(
        doc,
        ["层类型", "初始化规则", "目的"],
        [
            ["Conv2d", "Kaiming normal，fan_out，ReLU", "适应 ReLU，稳定深层信号传播"],
            ["BatchNorm weight", "全部设为1", "初始时不额外缩放归一化结果"],
            ["BatchNorm bias", "全部设为0", "初始时不额外平移归一化结果"],
            ["Linear", "保留 PyTorch 默认初始化", "满足当前设计，不增加额外规则"],
        ],
        [2500, 3600, 3260],
    )
    add_body(
        doc,
        "初始化在所有层创建完成后执行。测试既验证 _initialize_weights() 在构造模型时被调用，"
        "也遍历全部 BatchNorm 检查实际参数值，避免只存在方法却从未执行。"
    )

    doc.add_heading("八、测试先行的实践", level=1)
    add_code(
        doc,
        "RED：先写预期行为测试并观察失败\n"
        "GREEN：只写足够让测试通过的最小代码\n"
        "VERIFY：运行单项测试和完整测试，确认旧功能没有退化\n"
        "REPEAT：进入下一个更大的行为"
    )
    add_body(
        doc,
        "本阶段从“PlainBlock 类是否存在”开始，逐步扩展到第一组卷积、完整模块、参数校验、"
        "stem、三个阶段、完整 forward、初始化、统一模型入口、四种深度和 GPU 验证。"
        "这种方法让错误定位范围很小，也能明确每段代码为什么存在。"
    )
    add_labeled_paragraph(
        doc,
        "测试文件与模型文件的关系",
        "models/plainnet.py 负责实现功能；tests/test_models.py 导入正式模型、提供输入、调用功能并用 assert 判断结果。",
    )

    doc.add_heading("九、出现的问题与处理", level=1)
    add_table(
        doc,
        ["问题", "根本原因", "处理结果"],
        [
            ["完整模块仍返回原输入", "只添加了网络层，忘记修改 forward()", "连接六层数据流后测试通过"],
            ["pytest提示 no tests ran", "测试函数尚未写入或文件未保存", "确认函数名存在后重新运行"],
            ["非法参数只出现底层警告", "PyTorch允许先创建零元素层", "在 PlainBlock/PlainNet 构造阶段主动校验"],
            ["完整项目出现2项旧测试失败", "train.py 缺少 paths 且 resume 检查顺序退化", "恢复 paths，并优先检查参数冲突"],
            ["终端中文显示乱码", "终端显示编码问题", "十六进制确认源文件仍为正确 UTF-8"],
        ],
        [2600, 3600, 3160],
    )
    add_body(
        doc,
        "这次排错强调先读取错误栈、复现问题、定位具体文件和失败位置，再修改单一根因。"
        "最终完整项目测试恢复为37项全部通过。"
    )

    doc.add_heading("十、最终实验与验证结果", level=1)
    add_table(
        doc,
        ["模型", "真实深度", "可训练参数量", "CPU输出"],
        [
            ["PlainNet-20", "20", "269,722", "[2,10]"],
            ["PlainNet-32", "32", "464,154", "[2,10]"],
            ["PlainNet-44", "44", "658,586", "[2,10]"],
            ["PlainNet-56", "56", "853,018", "[2,10]"],
        ],
        [2500, 1900, 2660, 2300],
    )
    add_table(
        doc,
        ["验收项", "结果"],
        [
            ["完整自动测试", "37 passed，0 failed"],
            ["GPU型号", "NVIDIA GeForce RTX 3060 Laptop GPU"],
            ["GPU输入", "[1,3,32,32]，设备 cuda"],
            ["GPU输出", "[1,10]，设备 cuda:0"],
            ["训练状态", "未加载数据集、未反向传播、未保存训练结果"],
        ],
        [3000, 6360],
    )

    doc.add_heading("十一、参数量增长分析", level=1)
    add_body(
        doc,
        "相邻两种模型的参数量都增加194,432。因为 n 每次增加2，三个阶段各增加2个 PlainBlock，"
        "总共新增6个模块。16、32、64通道的单个 PlainBlock 参数量分别为4,672、18,560、73,984。"
    )
    add_code(
        doc,
        "每个阶段各增加1个模块：4,672 + 18,560 + 73,984 = 97,216\n"
        "n每次增加2：97,216 × 2 = 194,432"
    )
    add_body(
        doc,
        "实际参数量差值与理论计算完全一致，进一步证明三个阶段的模块数、通道数、卷积 bias 和 BatchNorm 参数均符合设计。"
    )

    doc.add_heading("十二、我在本阶段学到的内容", level=1)
    add_labeled_paragraph(
        doc,
        "模型结构",
        "理解了 Conv-BN-ReLU、stride、padding、通道数和特征图尺寸之间的关系。",
    )
    add_labeled_paragraph(
        doc,
        "模块化设计",
        "学会用 PlainBlock 表示可复用的基本单元，再通过 _make_stage() 组成深层网络。",
    )
    add_labeled_paragraph(
        doc,
        "前向传播",
        "理解 __init__() 负责创建层，forward() 才决定数据真正经过哪些层。",
    )
    add_labeled_paragraph(
        doc,
        "实验可信度",
        "知道不能只看模型名称或 depth 属性，而要检查真实层数、中间形状、参数量和最终输出。",
    )
    add_labeled_paragraph(
        doc,
        "测试与调试",
        "掌握了先写测试、观察失败、最小实现、完整回归，以及根据错误栈定位根因的基本流程。",
    )
    add_labeled_paragraph(
        doc,
        "GPU验证",
        "理解一次 CUDA 前向传播只能证明模型能在 GPU 上计算，不代表训练流程已经完成。",
    )

    doc.add_heading("十三、可直接用于日报的简要总结", level=1)
    add_body(
        doc,
        "今日完成 ResNet 论文复现项目第三阶段：基于 PyTorch 从零实现 CIFAR-10 版本的 "
        "PlainNet。首先实现不含残差连接的 PlainBlock，结构为两组 Conv-BN-ReLU；随后通过 "
        "_make_stage() 构建16、32、64通道的三个阶段，并加入全局平均池化和全连接分类器。"
        "完成 PlainNet-20/32/44/56 的统一模型注册、参数校验和 Kaiming 权重初始化。"
        "采用测试先行方式逐步验证模块形状、阶段降采样、真实网络深度、自定义类别数及模型入口，"
        "最终完整项目37项自动测试全部通过。四种模型参数量分别为269,722、464,154、658,586、"
        "853,018；PlainNet-20 在 RTX 3060 Laptop GPU 上完成前向传播，输出形状为[1,10]。"
        "本阶段只完成网络结构与计算验证，尚未加载 CIFAR-10 或执行训练。"
    )

    doc.add_heading("十四、下一阶段计划", level=1)
    add_labeled_paragraph(
        doc,
        "下一目标",
        "实现 CIFAR-10 ResNet-20/32/44/56，并与同深度 PlainNet 保持公平结构对比。",
    )
    add_labeled_paragraph(
        doc,
        "重点原理",
        "理解 ResidualBlock 中 F(x)+x 的作用，以及输入输出形状不同时的 Option A shortcut。",
    )
    add_labeled_paragraph(
        doc,
        "验收要求",
        "检查残差相加形状、Option A 无可训练参数、真实深度、中间尺寸、模型参数量与 GPU 前向传播。",
    )
    add_labeled_paragraph(
        doc,
        "后续训练",
        "PlainNet 与 ResNet 结构都通过后，再进入 CIFAR-10 数据、loss、优化器和5轮预实验。",
    )

    doc.add_heading("十五、后续追加区", level=1)
    add_table(
        doc,
        ["记录项", "填写内容"],
        [
            ["新问题", ""],
            ["我的理解", ""],
            ["运行命令", ""],
            ["观察结果", ""],
            ["问题原因", ""],
            ["解决方法", ""],
            ["下一步", ""],
        ],
        [2400, 6960],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
