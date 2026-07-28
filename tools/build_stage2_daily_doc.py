from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("outputs/2026-07-27 ResNet复现学习日报.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, bold=False, color="000000", name="Microsoft YaHei") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_code(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.05
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F6F8FA")
    p_pr.append(shd)
    run = p.add_run(text)
    set_run_font(run, size=9.5, name="Consolas")
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=10, bold=True, color=DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_run_font(p.add_run(value), size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_check_line(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    set_run_font(p.add_run("☐ "), size=11, color=BLUE)
    set_run_font(p.add_run(text), size=10.5)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    end_run = paragraph.add_run(" 页")
    set_run_font(end_run, size=9, color=MID_GRAY)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    style_specs = {
        "Title": (23, "000000", 0, 4),
        "Subtitle": (11, MID_GRAY, 0, 12),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in style_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("ResNet 论文复现 · 学习日报"), size=9, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("ResNet 复现学习日报"), size=23, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    set_run_font(
        subtitle.add_run("阶段二：实验骨架与配置系统总结"),
        size=12,
        color=MID_GRAY,
    )

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["日期", "2026年7月27日"],
            ["项目目录", r"E:\文档\resnet复现"],
            ["当前阶段", "阶段二完成；准备进入阶段三 PlainNet"],
            ["今日学习目标", "理解第二阶段代码，而不是只完成复制粘贴"],
        ],
        [1800, 7560],
    )

    doc.add_heading("一、今天完成了什么", level=1)
    add_body(
        doc,
        "第二阶段没有训练神经网络，而是建立一套可靠的实验管理骨架。"
        "它负责回答：训练哪个模型、使用什么参数、结果保存在哪里、"
        "采用短实验还是论文日程、是否从 checkpoint 恢复，以及参数是否合法。"
    )
    add_table(
        doc,
        ["工作", "当前作用", "为什么需要"],
        [
            ["项目目录分工", "分开模型、测试、数据和结果", "防止代码和实验产物混乱"],
            ["模型注册表", "统一记录 PlainNet/ResNet 型号", "通过名称创建模型"],
            ["命令行入口", "用 train.py 启动所有实验", "更换实验时无需修改源码"],
            ["参数校验", "训练前拦截不合理配置", "避免浪费训练时间"],
            ["实验命名与路径", "区分每一次运行", "防止结果互相覆盖"],
            ["YAML 配置", "记录完整实验条件", "保证结果可追溯、可复现"],
            ["恢复训练接口", "预留 checkpoint 续训能力", "长时间训练中断后可继续"],
            ["自动测试", "保护已经约定的规则", "后续改代码时及时发现退化"],
        ],
        [1800, 3100, 4460],
    )

    doc.add_heading("二、整个程序的执行流程", level=1)
    add_code(
        doc,
        "命令行参数 → 参数检查 → 查询模型配置 → 生成实验名称与路径\n"
        "→ 整理完整配置 → YAML往返验证 → 打印配置\n"
        "→（后续阶段）创建模型 → 加载数据 → 训练与测试 → 保存结果"
    )
    add_body(
        doc,
        "第二阶段完成到“打印配置”为止。因此执行 "
        "python train.py --model resnet20 时，不会下载数据、创建模型或开始训练，"
        "也不会生成 checkpoint 和 TensorBoard 日志。这是当前阶段的正常结果。"
    )

    doc.add_heading("三、文件分别负责什么", level=1)
    add_table(
        doc,
        ["文件/目录", "职责", "当前状态"],
        [
            ["train.py", "读取参数、规划实验、以后组织训练", "已完成配置骨架"],
            ["models/__init__.py", "模型注册表和统一创建入口", "接口已完成，模型未实现"],
            ["models/plainnet.py", "PlainNet 网络结构", "仅有职责说明"],
            ["models/resnet.py", "ResNet、残差块和 Option A", "仅有职责说明"],
            ["tests/test_train.py", "验证配置与参数规则", "已有2项测试"],
            ["tests/test_models.py", "验证模型形状、深度和公平性", "目前记录验收目标"],
            ["outputs", "checkpoint、CSV、实验配置", "尚未开始写入"],
            ["runs", "TensorBoard 事件日志", "尚未开始写入"],
        ],
        [2100, 4200, 3060],
    )

    doc.add_heading("四、模型注册表：为什么先登记模型", level=1)
    add_body(
        doc,
        "MODEL_CONFIGS 是模型“目录”。它把模型名称映射到 family、depth 和 n。"
        "family 表示 PlainNet 或 ResNet；depth 是论文中的网络深度；n 是 CIFAR 网络"
        "每个阶段重复的基本模块数量。"
    )
    add_table(
        doc,
        ["模型", "n", "深度公式", "结果"],
        [
            ["PlainNet/ResNet-20", "3", "6×3+2", "20层"],
            ["PlainNet/ResNet-32", "5", "6×5+2", "32层"],
            ["PlainNet/ResNet-44", "7", "6×7+2", "44层"],
            ["PlainNet/ResNet-56", "9", "6×9+2", "56层"],
        ],
        [3000, 1100, 2700, 2560],
    )
    add_body(
        doc,
        "深度公式来自：开头1个卷积层 + 3个阶段×每阶段n个模块×每模块2个卷积层 "
        "+ 最后1个全连接层，即 1+3×n×2+1=6n+2。PlainNet 与 ResNet 使用相同深度，"
        "才能把主要变量控制为“是否存在残差连接”。"
    )

    doc.add_heading("五、create_model()：统一模型入口", level=1)
    add_code(doc, "def create_model(name: str, num_classes: int = 10) -> nn.Module:")
    add_body(
        doc,
        "name 是模型名称；num_classes 是最终输出类别数。CIFAR-10 有10类，所以默认值为10。"
        "若输入一批2张图片，最终分类输出应为 [2, 10]：2代表样本数，10代表每张图片的10个类别分数。"
    )
    add_body(
        doc,
        "函数会检查 name 是否为字符串、名称是否在注册表中、num_classes 是否为正整数。"
        "当前合法模型最终仍抛出 NotImplementedError，这是明确表示“接口已设计，但网络结构尚未完成”。"
        "第三、第四阶段完成后，这里才会真正返回 PlainNet 或 ResNet 对象。"
    )

    doc.add_heading("六、parse_args()：把终端命令变成实验参数", level=1)
    add_code(doc, "python train.py --model resnet20 --epochs 5")
    add_body(
        doc,
        "argparse 会把终端文字解析为 args.model='resnet20'、args.epochs=5 等属性。"
        "args 可以理解成一个参数收纳盒。未在命令中写出的参数使用默认值。"
    )
    add_table(
        doc,
        ["参数", "默认值", "作用"],
        [
            ["--model", "必须指定", "选择 plain20～56 或 resnet20～56"],
            ["--epochs", "5", "短实验训练轮数"],
            ["--batch-size", "128", "一次送入网络的图片数"],
            ["--learning-rate", "0.1", "SGD 初始学习率"],
            ["--momentum", "0.9", "利用历史更新方向"],
            ["--weight-decay", "0.0001", "权重衰减/正则化"],
            ["--seed", "42", "控制初始化和数据随机过程"],
            ["--num-workers", "0", "Windows 初期由主进程加载数据"],
            ["--run-name", "None", "可选的自定义实验名称"],
            ["--device", "cuda", "使用 GPU 或 CPU"],
            ["--schedule", "short", "短实验或论文正式日程"],
            ["--max-iterations", "64000", "论文日程最大参数更新次数"],
            ["--resume", "None", "从指定 .pt checkpoint 继续"],
        ],
        [2700, 1700, 4960],
    )

    doc.add_heading("七、epoch、batch 与 iteration", level=1)
    add_body(
        doc,
        "epoch 表示模型完整看过一次训练集；batch 是一次送入网络的一组图片；"
        "iteration 表示用一个 batch 完成一次前向传播、反向传播和参数更新。"
    )
    add_code(
        doc,
        "CIFAR-10训练集：50000张图片\n"
        "batch_size=128 → 每个epoch约 ceil(50000/128)=391次iteration\n"
        "论文64000 iterations → 大约163.7个epoch"
    )
    add_body(
        doc,
        "因此 short 日程按 epoch 控制，适合先用5轮确认流程有效；paper 日程按 iteration 控制，"
        "并在32000和48000次更新时调整学习率。5轮实验只能验证流程和趋势，不能作为论文正式结果。"
    )

    doc.add_heading("八、validate_args()：训练前主动发现错误", level=1)
    add_body(
        doc,
        "argparse 能确认参数类型，却不知道参数是否合理。例如 -5 是整数，但不能训练负5轮。"
        "validate_args() 因此继续检查数值范围、实验名称、论文日程和 checkpoint。"
    )
    add_table(
        doc,
        ["检查对象", "规则", "目的"],
        [
            ["epochs/batch size/学习率", "必须大于0", "排除无意义训练"],
            ["momentum", "0≤momentum<1", "限制为合理范围"],
            ["weight decay/seed/workers", "不得小于0", "排除非法设置"],
            ["run name", "非空且不能包含路径", "防止写入意外目录"],
            ["paper iterations", "必须超过48000", "确保经过第二个学习率节点"],
            ["resume + run name", "不能同时出现", "避免续训和新实验语义冲突"],
            ["resume 文件", "存在、是文件、扩展名为.pt", "训练前确认 checkpoint 可用"],
        ],
        [2700, 3100, 3560],
    )

    doc.add_heading("九、实验名称与保存路径", level=1)
    add_body(
        doc,
        "未指定 --run-name 时，短实验自动命名为 seed42_5epochs；论文日程命名为 "
        "seed42_paper64k。名称包含关键条件，方便快速识别实验。"
    )
    add_code(
        doc,
        "outputs/resnet20/seed42_5epochs   → checkpoint、CSV、配置\n"
        "runs/resnet20/seed42_5epochs      → TensorBoard日志"
    )
    add_body(
        doc,
        "build_experiment_paths() 当前只计算路径，没有调用 mkdir，所以测试配置时不会留下空目录。"
        "真正开始训练时才会创建目录。如果使用 --resume，输出目录取 checkpoint 的父目录，"
        "从而继续写入原实验，而不是另建一个含义不清的新实验。"
    )

    doc.add_heading("十、实验配置与 YAML 往返验证", level=1)
    add_body(
        doc,
        "build_experiment_config() 把分散的命令行参数整理为 model、data、training、paths 四组信息。"
        "以后保存为 config.yml 后，即使隔了几个月，也能还原模型、数据处理、优化器、随机种子、"
        "训练长度、恢复来源和输出位置。"
    )
    add_code(
        doc,
        "Python字典 → YAML文本 → 再读取成Python字典 → 与原字典比较"
    )
    add_body(
        doc,
        "validate_config_round_trip() 执行上述往返。如果读回内容与原配置不同，程序立即报错。"
        "这保证配置不是“看起来保存了”，实际却发生字段丢失或类型变化。"
    )

    doc.add_heading("十一、main() 如何串起第二阶段", level=1)
    add_table(
        doc,
        ["顺序", "函数", "结果"],
        [
            ["1", "parse_args()", "读取命令行并完成参数校验"],
            ["2", "build_experiment_paths()", "计算输出与日志目录"],
            ["3", "build_experiment_config()", "整理完整实验配置"],
            ["4", "validate_config_round_trip()", "验证 YAML 可无损保存"],
            ["5", "print()", "打印原始参数、路径和整理后的配置"],
        ],
        [900, 3600, 4860],
    )
    add_body(
        doc,
        "if __name__ == '__main__': main() 表示只有直接运行 train.py 时才执行主程序。"
        "测试代码导入 train.py 中的函数时不会意外启动训练，这使各个函数可以被安全复用和测试。"
    )

    doc.add_heading("十二、自动测试在保护什么", level=1)
    add_body(
        doc,
        "tests/test_train.py 使用一组假的参数检查实验配置逻辑，不会启动真实训练。"
        "当前两个测试分别确认：配置正确记录 resume_from 和 paths；当 resume 与 run_name 同时出现时，"
        "优先报告二者冲突，而不是先报告 checkpoint 不存在。"
    )
    add_body(
        doc,
        "pytest.ini 中的 -p no:cacheprovider 只是关闭 pytest 缓存插件，解决旧 .pytest_cache 的"
        "Windows 权限警告。它不会关闭测试，也不会影响模型或训练结果。当前验证结果为：2 passed。"
    )

    doc.add_heading("十三、第二阶段应真正掌握的思想", level=1)
    insights = [
        "实验参数不应散落在源码中：使用命令行可以在不改训练代码的情况下切换模型和超参数。",
        "实验必须可追溯：准确率必须与模型、数据处理、优化器、随机种子和训练长度一起保存。",
        "错误要尽早发现：明显错误应在占用 GPU 之前终止。",
        "模型结构与训练流程要分离：PlainNet 和 ResNet 才能共用同一训练流程并公平比较。",
        "测试是在保护实验规则：后续修改功能时，已有约定不能被悄悄破坏。",
    ]
    for item in insights:
        add_check_line(doc, item)

    doc.add_heading("十四、今日完成情况与验证记录", level=1)
    add_table(
        doc,
        ["检查项", "状态", "说明"],
        [
            ["Conda环境与GPU", "已通过", "resnet-paper；RTX 3060 Laptop；CUDA可用"],
            ["训练命令参数解析", "已通过", "resnet20默认配置可正常显示"],
            ["实验路径生成", "已通过", "outputs与runs按模型/实验名分层"],
            ["YAML配置整理", "已通过", "模型、数据、训练、路径字段齐全"],
            ["配置测试", "已通过", "pytest：2 passed，0 warnings"],
            ["真实模型与训练", "未开始", "将在后续阶段实现"],
        ],
        [3200, 1700, 4460],
    )

    doc.add_heading("十五、今日问题与理解记录（持续追加）", level=1)
    add_table(
        doc,
        ["时间", "问题/困惑", "当前理解或处理结果"],
        [
            ["今日", "为什么目前没有生成训练文件？", "第二阶段只规划路径和配置，真实训练时才创建目录并保存文件。"],
            ["今日", "num_classes=10 是什么？", "代表每张图片最终输出10个类别分数，不是图片数量或卷积尺寸。"],
            ["今日", "为什么需要自动测试？", "用自动规则保护实验接口，防止后续修改造成旧功能退化。"],
            ["待追加", "", ""],
            ["待追加", "", ""],
        ],
        [1100, 3600, 4660],
    )

    doc.add_heading("十六、下一步计划", level=1)
    add_body(
        doc,
        "下一阶段进入 PlainNet 基线实现。学习顺序应坚持“先理解，再写代码，再验证”："
    )
    next_steps = [
        "理解 CIFAR 输入张量 [N, 3, 32, 32] 中每个维度的含义。",
        "理解卷积核、通道数、stride、padding 对输出形状的影响。",
        "设计并实现最小 PlainBlock：卷积、批归一化和 ReLU，不包含残差连接。",
        "用随机张量验证 PlainBlock 的输入输出形状。",
        "搭建三个阶段的 PlainNet，并验证 20/32/44/56 层深度。",
        "将模型接入 create_model()，补充模型结构测试。",
    ]
    for step in next_steps:
        add_check_line(doc, step)

    doc.add_heading("十七、后续日报追加模板", level=1)
    add_body(doc, "后续每完成一个小步骤，就在这里追加记录，避免只保存最终结果而丢失学习过程。")
    add_table(
        doc,
        ["记录项", "填写内容"],
        [
            ["完成的代码", "文件、函数或模块名称："],
            ["我理解的原理", "用自己的话说明："],
            ["运行命令", ""],
            ["观察到的输出", ""],
            ["遇到的问题", ""],
            ["问题原因", ""],
            ["解决方法", ""],
            ["验证结果", ""],
            ["下一步", ""],
        ],
        [2400, 6960],
    )

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.line_spacing = 1.2
    set_run_font(note.add_run("学习原则："), bold=True, color=DARK_BLUE)
    set_run_font(
        note.add_run(
            "每次复制代码前，先说明它接收什么、输出什么、为什么存在；"
            "复制后必须亲自运行，并根据结果解释程序实际做了什么。"
        ),
        color=DARK_BLUE,
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
